"""Отчёты: Excel-ведомость расхождений и короткая сводка для Telegram."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Document

FILL = {"FAIL": "F8CBAD", "WARNING": "FFE699", "NEEDS_CONTEXT": "BDD7EE", "INFO": "E7E6E6", "OK": "C6EFCE"}
ICON = {"FAIL": "🔴", "WARNING": "🟡", "NEEDS_CONTEXT": "🔵", "INFO": "⚪", "OK": "🟢"}
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def price_factor(lsr: Document) -> float:
    t = lsr.totals
    if "grand" in t and "vsego" in t and t["vsego"][0]:
        return t["grand"][0] / t["vsego"][0]
    return 1.0


def _head(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row, i, h)
        c.font, c.alignment, c.border = FB, Alignment(wrap_text=True, vertical="center"), THIN
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)


def to_xlsx(path: str, lsr: Document, acts: list, findings: list, table: list) -> None:
    wb = Workbook()

    # --- Замечания
    wz = wb.active
    wz.title = "Замечания"
    hdr = ["№", "Правило", "Статус", "Замечание", "Документы", "Где", "Пояснение",
           "Возможное завышение, руб. (до коэф.)"]
    _head(wz, 1, hdr, [5, 9, 15, 38, 34, 26, 90, 18])
    for i, f in enumerate(findings, 1):
        vals = [i, f.rule_id, f.status, f.title, f.docs, f.location, f.explanation, f.delta_rub]
        for j, v in enumerate(vals, 1):
            c = wz.cell(i + 1, j, v)
            c.font, c.alignment, c.border = F, WRAP, THIN
        wz.cell(i + 1, 3).fill = PatternFill("solid", fgColor=FILL[f.status])
        wz.cell(i + 1, 8).number_format = '#,##0.00;-#,##0.00;"-"'
    wz.auto_filter.ref = f"A1:H{len(findings) + 1}"
    last = len(findings) + 1

    # --- Сводка
    ws = wb.create_sheet("Сводка", 0)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 22
    ws["A1"] = "Сверка ЛСР ↔ КС-2 ↔ КС-3 (этап 1, детерминированные проверки)"
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    rows = [("Объект", lsr.meta.get("объект", "")[:120]),
            ("Смета", f"ЛСР № {lsr.meta.get('lsr_number', '')}, {lsr.meta.get('price_level', '')}"),
            ("Акты КС-2", ", ".join(f"№{a.meta.get('act_number')} от {a.meta.get('act_date')}" for a in acts)),
            ("Контракт", f"{acts[0].meta.get('contract_number', '')} от {acts[0].meta.get('contract_date', '')}")]
    r = 3
    for k, v in rows:
        ws.cell(r, 1, k).font = FB
        ws.cell(r, 2, v).font = F
        r += 1
    r += 1
    ws.cell(r, 1, "Статус").font = FB
    ws.cell(r, 2, "Количество").font = FB
    for st in ("FAIL", "WARNING", "NEEDS_CONTEXT", "INFO", "OK"):
        r += 1
        ws.cell(r, 1, st).fill = PatternFill("solid", fgColor=FILL[st])
        ws.cell(r, 1).font = F
        ws.cell(r, 2, f'=COUNTIF(Замечания!$C$2:$C${last},A{r})').font = F
    r += 2
    ws.cell(r, 1, "Возможное завышение, прямые затраты (до коэф.), руб.").font = FB
    ws.cell(r, 2, f'=SUMIFS(Замечания!$H$2:$H${last},Замечания!$C$2:$C${last},"<>OK",Замечания!$H$2:$H${last},">0")').font = F
    ws.cell(r, 2).number_format = "#,##0.00"
    r1 = r
    r += 1
    ws.cell(r, 1, "Коэффициент перехода к цене контракта (ВСЕГО / Всего до снижения)").font = F
    ws.cell(r, 2, round(price_factor(lsr), 6)).font = Font(name="Arial", size=10, color="0000FF")
    ws.cell(r, 2).number_format = "0.000000"
    ws.cell(r, 3, f"Источник: ЛСР, {lsr.totals.get('grand', ('',))[0]} / {lsr.totals.get('vsego', ('',))[0]} "
                  "(коэф. снижения × надбавка УСН)").font = Font(name="Arial", size=9, italic=True)
    r += 1
    ws.cell(r, 1, "То же в ценах контракта, руб. (оценка)").font = FB
    ws.cell(r, 2, f"=B{r1}*B{r - 1}").font = FB
    ws.cell(r, 2).number_format = "#,##0.00"
    r += 1
    ws.cell(r, 1, "Недооценено в смете (материалов меньше нормы), прямые затраты, руб.").font = F
    ws.cell(r, 2, f'=-SUMIFS(Замечания!$H$2:$H${last},Замечания!$C$2:$C${last},"<>OK",Замечания!$H$2:$H${last},"<0")').font = F
    ws.cell(r, 2).number_format = "#,##0.00"
    r += 2
    ws.cell(r, 1, "Ограничения: проверяется внутренняя непротиворечивость документов. Фактическое выполнение "
                  "объёмов не подтверждается (нужна сверка с АОСР/ИД). Правила H-* — эвристики, "
                  "требуют подтверждения инженером.").font = Font(name="Arial", size=9, italic=True)
    ws.cell(r, 1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 40

    # --- Сопоставление
    wc = wb.create_sheet("Сопоставление")
    h = ["№ поз.", "Шифр", "Наименование", "Ед.", "Объём ЛСР", "Объём КС-2 нараст.", "% закрытия",
         "Цена ед. ЛСР", "Цена ед. КС-2", "Δ цены", "Сумма ЛСР", "Сумма КС-2", "Δ суммы"]
    _head(wc, 1, h, [8, 24, 60, 9, 13, 13, 10, 13, 13, 11, 14, 14, 12])
    for i, t in enumerate(table, 2):
        vals = [t["key"], t["code"], ("    " if t["is_sub"] else "") + t["name"], t["unit"],
                t["lsr_qty"], t["ks_qty"], f'=IF(E{i}=0,"-",F{i}/E{i})',
                t["lsr_price"], t["ks_price"], f'=IF(OR(H{i}="",I{i}=""),"",I{i}-H{i})',
                t["lsr_total"], t["ks_total"], f'=IF(K{i}="","",L{i}-K{i})']
        for j, v in enumerate(vals, 1):
            c = wc.cell(i, j, v)
            c.font, c.border = (Font(name="Arial", size=9, color="595959") if t["is_sub"] else F), THIN
        for j in (5, 6):
            wc.cell(i, j).number_format = "#,##0.0000"
        wc.cell(i, 7).number_format = "0.0%"
        for j in (8, 9, 10, 11, 12, 13):
            wc.cell(i, j).number_format = '#,##0.00;-#,##0.00;"-"'
    wc.auto_filter.ref = f"A1:M{len(table) + 1}"
    wb.save(path)


def to_telegram(lsr: Document, acts: list, findings: list) -> str:
    cnt = {s: sum(f.status == s for f in findings) for s in ICON}
    over = sum(f.delta_rub for f in findings if f.status != "OK" and (f.delta_rub or 0) > 0)
    under = -sum(f.delta_rub for f in findings if f.status != "OK" and (f.delta_rub or 0) < 0)
    k = price_factor(lsr)
    lines = [f"Сверка ЛСР № {lsr.meta.get('lsr_number', '')} ↔ КС-2 ({len(acts)} шт.)",
             " · ".join(f"{ICON[s]} {cnt[s]}" for s in ("FAIL", "WARNING", "NEEDS_CONTEXT", "OK")),
             ""]
    if over:
        lines.append(f"Возможное завышение: {over:,.0f} ₽ прямых затрат ≈ {over * k:,.0f} ₽ в ценах контракта"
                     .replace(",", " "))
        if under:
            lines.append(f"Недооценено (не в пользу подрядчика): {under:,.0f} ₽".replace(",", " "))
        lines.append("")
    for f in [f for f in findings if f.status in ("FAIL", "WARNING")][:6]:
        rub = f" — {f.delta_rub:,.0f} ₽".replace(",", " ") if f.delta_rub else ""
        lines.append(f"{ICON[f.status]} {f.rule_id} {f.title}: {f.location}{rub}")
    nc = list(dict.fromkeys(f.title for f in findings if f.status == "NEEDS_CONTEXT"))
    if nc:
        lines += ["", "Уточнить: " + "; ".join(nc)]
    lines += ["", "Полный отчёт — в Excel-файле."]
    return "\n".join(lines)
