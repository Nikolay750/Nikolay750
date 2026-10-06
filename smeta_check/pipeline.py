"""Единая точка входа для сервера: классифицировать загруженные файлы и запустить подходящие проверки.

Возвращает JSON-совместимый результат (для Mini App) и пишет Excel-отчёт.
"""
import os
import re

import openpyxl
import pdfplumber

from .models import Finding

ORDER = ("FAIL", "WARNING", "NEEDS_CONTEXT", "INFO", "OK")
TYPE_LABEL = {
    "lsr_xlsx": "Смета (ЛСР)", "ks2_xlsx": "Акт КС-2", "smeta_contract_xlsx": "Смета контракта",
    "contract_pdf": "Договор с приложениями", "scan_pdf": "Скан (без текста)",
    "pdf_other": "PDF", "xls_legacy": "Excel 97-2003 (.xls)", "unknown": "Не распознан",
}


def _pdf_text(path: str, pages: int = 0) -> str:
    """Step 9.1: раньше вызывал внешнюю программу `pdftotext` (poppler) через
    subprocess — на Windows её нет и не ставится через pip, из-за чего
    classify() падал с FileNotFoundError на любом PDF. pdfplumber — чистая
    Python-библиотека (ставится через pip install -r requirements.txt),
    layout=True старается сохранить колонки так же, как pdftotext -layout."""
    try:
        with pdfplumber.open(path) as pdf:
            pages_iter = pdf.pages[:pages] if pages else pdf.pages
            return "\n".join(p.extract_text(layout=True) or "" for p in pages_iter)
    except Exception:
        # Битый/зашифрованный/нетипичный PDF — не роняем всю проверку пакета,
        # classify() трактует пустой текст как скан (scan_pdf), это штатный путь.
        return ""


def classify(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        return "xls_legacy"
    if ext in (".xlsx", ".xlsm"):
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                text = " ".join(str(v) for row in wb.worksheets[0].iter_rows(max_row=80, values_only=True)
                                for v in row if v)
            finally:
                wb.close()          # read_only держит файл открытым — на сервере это утечка дескрипторов
        except Exception:
            return "unknown"
        t = re.sub(r"\s+", " ", text.upper())
        if "О ПРИЕМКЕ ВЫПОЛНЕННЫХ РАБОТ" in t or "О ПРИЁМКЕ ВЫПОЛНЕННЫХ РАБОТ" in t:
            return "ks2_xlsx"
        if "СМЕТЫ КОНТРАКТА" in t or "СМЕТА КОНТРАКТА" in t:
            return "smeta_contract_xlsx"
        if "ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЕТ" in t or "ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЁТ" in t:
            return "lsr_xlsx"
        return "unknown"
    if ext == ".pdf":
        head = _pdf_text(path, 3)
        if len(re.sub(r"\s", "", head)) < 200:
            return "scan_pdf"
        full = _pdf_text(path)
        if re.search(r"Договор|Контракт", full) and "Ведомость объ" in full and "Всего по позиции" in full:
            return "contract_pdf"
        return "pdf_other"
    return "unknown"


def _f2d(f: Finding) -> dict:
    return {"rule": f.rule_id, "status": f.status, "title": f.title, "docs": f.docs,
            "where": f.location, "text": f.explanation, "rub": round(f.delta_rub, 2) if f.delta_rub else None}


def check_estimate_acts(lsr_path, ks2_paths, out_xlsx):
    from .parser_grandsmeta import parse
    from .report import to_xlsx, price_factor
    from .rules import run_all
    lsr = parse(lsr_path, "ЛСР")
    lsr.doc_id = f"ЛСР {lsr.meta.get('lsr_number', '')}".strip()
    acts = []
    for p in ks2_paths:
        d = parse(p, "КС-2")
        d.doc_id = f"КС-2 №{d.meta.get('act_number', '?')}"
        acts.append(d)
    findings, table = run_all(lsr, acts)
    to_xlsx(out_xlsx, lsr, acts, findings, table)
    title = f"Смета № {lsr.meta.get('lsr_number', '')} ↔ КС-2 ({len(acts)} шт.)"
    return title, findings, None, price_factor(lsr)


def check_contract(pdf_path, smeta_xlsx, out_xlsx):
    from .models import Finding as Fd
    from .parser_pdf import parse_contract_price, parse_hidden_works, parse_lsr_pdf, parse_vor_pdf, pdf_pages
    from .report_contract import to_xlsx
    from .rules import _match_norm_materials, h01_diameters, h02_debris_balance, h03_replacement_qty, unit_mult
    from .rules_contract import c_price_vat, l_arithmetic, v_vor_lsr, z_hidden_works
    pages = pdf_pages(pdf_path)
    contract = parse_contract_price(pages)
    lsr = parse_lsr_pdf(pages, "ЛСР (прил. 7)")
    vor = parse_vor_pdf(pages)
    hw = parse_hidden_works(pages)
    sc = {}
    if smeta_xlsx:
        ws = openpyxl.load_workbook(smeta_xlsx, data_only=True).active
        for row in ws.iter_rows(values_only=True):
            label = " ".join(str(x) for x in row if isinstance(x, str))
            nums = [x for x in row if isinstance(x, (int, float))]
            if nums and re.search(r"цена контракта с НДС|^\s*Итого", label):
                sc["total"] = nums[-1]
    findings = c_price_vat(contract, lsr, sc) + l_arithmetic(lsr)
    f, vor_rows = v_vor_lsr(vor, lsr)
    findings += f
    f, plan = z_hidden_works(hw, lsr, vor)
    findings += f
    findings += h01_diameters(lsr) + h02_debris_balance(lsr) + h03_replacement_qty(lsr)
    n_norm = n_cmp = 0
    for p in lsr.positions.values():
        if p.is_block:
            for r, mats, _ in _match_norm_materials(lsr, p):
                n_norm += 1
                n_cmp += bool(mats and unit_mult(r.unit)[1] == unit_mult(mats[0].unit)[1])
    if not any(x.rule_id == "H-03" for x in findings):
        findings.append(Fd("H-03", "OK", "Количество материалов сверено с нормами ГЭСН", lsr.ref(), "—",
                           f"Строк неучтённых материалов в нормах: {n_norm}, сверено по количеству: {n_cmp}. "
                           "Отклонений больше ±33% не найдено."))
    findings.sort(key=lambda x: (ORDER.index(x.status), -abs(x.delta_rub or 0), x.rule_id))
    title = f"Договор № {contract.get('number', '')}: договор ↔ ЛСР ↔ ВОР ↔ ТЗ"
    info = [("Договор", f"№ {contract.get('number')}"), ("ЛСР", f"{len(lsr.positions)} позиций"),
            ("ВОР", f"{len(vor)} позиций"), ("Перечень скрытых работ", f"{len(hw)} работ")]
    to_xlsx(out_xlsx, title, info, findings, plan, vor_rows)
    return title, findings, plan, None


def run_package(paths: list, out_dir: str) -> dict:
    files = [{"name": os.path.basename(p), "path": p, "type": classify(p), "size": os.path.getsize(p)}
             for p in paths]
    by = lambda t: [f["path"] for f in files if f["type"] == t]
    extra = []
    for f in files:
        if f["type"] == "scan_pdf":
            extra.append(Finding("IN-01", "NEEDS_CONTEXT", "Скан без текстового слоя", f["name"], "—",
                                 "Распознавание сканов (OCR) в этой версии не выполняется. Если есть исходник "
                                 "(xlsx из сметной программы или PDF, сохранённый из программы), загрузите его."))
        elif f["type"] == "xls_legacy":
            extra.append(Finding("IN-01", "NEEDS_CONTEXT", "Файл в старом формате .xls", f["name"], "—",
                                 "Пересохраните в .xlsx (Файл → Сохранить как → Книга Excel) и загрузите снова."))
        elif f["type"] in ("unknown", "pdf_other"):
            extra.append(Finding("IN-01", "INFO", "Файл не участвовал в проверке", f["name"], "—",
                                 "Тип документа не распознан: ожидаются ЛСР/КС-2 в xlsx или договор с ВОР и ЛСР "
                                 "в PDF с текстовым слоем."))
    out_xlsx = os.path.join(out_dir, "report.xlsx")
    title, findings, plan, factor, mode = "Пакет без проверяемых пар", [], None, None, "none"
    if by("contract_pdf"):
        title, findings, plan, factor = check_contract(by("contract_pdf")[0],
                                                       (by("smeta_contract_xlsx") or [None])[0], out_xlsx)
        mode = "contract"
    elif by("lsr_xlsx") and by("ks2_xlsx"):
        title, findings, plan, factor = check_estimate_acts(by("lsr_xlsx")[0], by("ks2_xlsx"), out_xlsx)
        mode = "acts"
    elif by("lsr_xlsx"):
        extra.append(Finding("IN-02", "NEEDS_CONTEXT", "Смета загружена без актов КС-2", "—", "—",
                             "Для сверки нужны акты КС-2 к этой смете (xlsx). Проверки самой сметы появятся позже."))
    elif by("ks2_xlsx"):
        extra.append(Finding("IN-02", "NEEDS_CONTEXT", "Акты КС-2 загружены без сметы", "—", "—",
                             "Для сверки нужна ЛСР, по которой составлены акты (xlsx)."))
    allf = findings + extra
    allf.sort(key=lambda x: (ORDER.index(x.status), -abs(x.delta_rub or 0), x.rule_id))
    over = sum(f.delta_rub for f in allf if f.status != "OK" and (f.delta_rub or 0) > 0 and f.rule_id.startswith("H"))
    under = -sum(f.delta_rub for f in allf if f.status != "OK" and (f.delta_rub or 0) < 0)
    return {
        "mode": mode, "title": title,
        "files": [{"name": f["name"], "type": f["type"], "label": TYPE_LABEL[f["type"]], "size": f["size"]}
                  for f in files],
        "counts": {s: sum(f.status == s for f in allf) for s in ORDER},
        "over_rub": round(over, 2), "over_contract_rub": round(over * factor, 2) if factor else None,
        "under_rub": round(under, 2),
        "findings": [_f2d(f) for f in allf],
        "plan": plan, "report": "report.xlsx" if mode != "none" else None,
    }
