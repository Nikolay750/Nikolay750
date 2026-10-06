"""Отчёт по пакету договора (без КС-2): Excel + сводка для Telegram."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .report import FILL, ICON, F, FB, THIN, WRAP, _head


def to_xlsx(path, title, info, findings, plan, vor_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 70
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    r = 3
    for k, v in info:
        ws.cell(r, 1, k).font = FB
        c = ws.cell(r, 2, v)
        c.font, c.alignment = F, WRAP
        r += 1
    last = len(findings) + 1
    r += 1
    ws.cell(r, 1, "Статус").font = FB
    ws.cell(r, 2, "Количество").font = FB
    for st in ("FAIL", "WARNING", "NEEDS_CONTEXT", "INFO", "OK"):
        r += 1
        ws.cell(r, 1, st).fill = PatternFill("solid", fgColor=FILL[st])
        ws.cell(r, 1).font = F
        ws.cell(r, 2, f"=COUNTIF(Замечания!$C$2:$C${last},A{r})").font = F
        ws.cell(r, 2).alignment = Alignment(horizontal="left")

    wz = wb.create_sheet("Замечания")
    _head(wz, 1, ["№", "Правило", "Статус", "Замечание", "Документы", "Где", "Пояснение", "Сумма, руб."],
          [5, 9, 15, 40, 30, 26, 100, 15])
    for i, f in enumerate(findings, 1):
        for j, v in enumerate([i, f.rule_id, f.status, f.title, f.docs, f.location, f.explanation, f.delta_rub], 1):
            c = wz.cell(i + 1, j, v)
            c.font, c.alignment, c.border = F, WRAP, THIN
        wz.cell(i + 1, 3).fill = PatternFill("solid", fgColor=FILL[f.status])
        wz.cell(i + 1, 8).number_format = '#,##0.00;-#,##0.00;"-"'
    wz.auto_filter.ref = f"A1:H{last}"

    wp = wb.create_sheet("План АОСР")
    _head(wp, 1, ["№ АОСР по перечню ТЗ", "Скрытая работа (ТЗ прил. 9)", "Поз. ЛСР", "Раздел (помещение)",
                  "Объём по ВОР", "Ед.", "Совпадение"], [10, 70, 9, 32, 12, 8, 11])
    for i, x in enumerate(plan, 2):
        sec = x["section"].split(".", 1)[1].strip() if "." in x["section"] else x["section"]
        for j, v in enumerate([x["n"], x["name"], x["lsr"], sec, x["qty"], x["unit"], x["score"]], 1):
            c = wp.cell(i, j, v)
            c.font, c.border = F, THIN
        wp.cell(i, 7).number_format = "0%"
    wp.auto_filter.ref = f"A1:G{len(plan) + 1}"

    wv = wb.create_sheet("ВОР ↔ ЛСР")
    _head(wv, 1, ["п. ВОР", "поз. ЛСР", "Наименование (ВОР)", "Ед. ВОР", "Объём ВОР", "Формула ВОР",
                  "Ед. ЛСР", "Объём ЛСР", "Множитель", "ЛСР в ед. ВОР", "Δ", "Статус", "Примечание"],
          [7, 8, 60, 10, 11, 18, 10, 11, 10, 12, 10, 11, 20])
    for i, x in enumerate(vor_rows, 2):
        vals = [x["vor"], x["lsr"], x["name_vor"], x["unit_vor"], x["qty_vor"], x["formula"], x["unit_lsr"],
                x["qty_lsr"], x["mult"], f'=IF(OR(H{i}="",I{i}=""),"",H{i}*I{i})', f'=IF(J{i}="","",J{i}-E{i})',
                x["status"], x["note"]]
        for j, v in enumerate(vals, 1):
            c = wv.cell(i, j, v)
            c.font, c.border = F, THIN
        if x["status"] != "OK":
            wv.cell(i, 12).fill = PatternFill("solid", fgColor=FILL[x["status"]])
        wv.cell(i, 11).number_format = '0.####;-0.####;"-"'
    wv.auto_filter.ref = f"A1:M{len(vor_rows) + 1}"
    wb.save(path)


def to_telegram(title, findings):
    cnt = {s: sum(f.status == s for f in findings) for s in ICON}
    lines = [title, " · ".join(f"{ICON[s]} {cnt[s]}" for s in ("FAIL", "WARNING", "NEEDS_CONTEXT", "OK")), ""]
    for f in [f for f in findings if f.status in ("FAIL", "WARNING")][:6]:
        rub = f" — {abs(f.delta_rub):,.0f} ₽".replace(",", " ") if f.delta_rub else ""
        lines.append(f"{ICON[f.status]} {f.rule_id} {f.title}{rub}")
    lines += ["", "Полный отчёт и план АОСР — в Excel-файле."]
    return "\n".join(lines)
