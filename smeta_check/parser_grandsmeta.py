"""Парсер выгрузок ГРАНД-Сметы (ЛСР по форме 421/пр и КС-2 «по Методике»).

Колонки определяются по тексту шапки, а не по буквам: в ЛСР и КС-2 они сдвинуты,
а у других сметных программ порядок будет свой.
"""
import re
from typing import Optional

import openpyxl

from .models import Document, Overhead, Position, Resource

KEY_RE = re.compile(r"^\d+(\.\d+)*$")
MARKERS = {"Н", "П", "П,Н", "Уд", "О"}


def num(v) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def txt(v) -> str:
    return "" if v is None else str(v).strip()


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.lower()).strip()


def _find_header(ws):
    for r in range(1, min(ws.max_row, 150) + 1):
        texts = [_norm(txt(c.value)) for c in ws[r]]
        if any(t.startswith("обоснование") for t in texts) and any("наименование работ" in t for t in texts):
            return r
    raise ValueError("Не найдена шапка таблицы (строка с «Обоснование» и «Наименование работ»)")


def _map_columns(ws, hdr: int) -> tuple[dict, int]:
    # строка нумерации граф 1,2,3…
    num_row = None
    for r in range(hdr + 1, hdr + 6):
        vals = [num(c.value) for c in ws[r]]
        ints = [v for v in vals if v is not None]
        if len(ints) >= 8 and ints[:3] == [1.0, 2.0, 3.0]:
            num_row = r
            break
    if num_row is None:
        raise ValueError("Не найдена строка нумерации граф под шапкой")

    col_text = {}
    for r in range(hdr, num_row):
        for c in ws[r]:
            t = _norm(txt(c.value))
            if t:
                col_text.setdefault(c.column, []).append(t)
    cols = {}
    qty_cols = []
    for col, parts in sorted(col_text.items()):
        t = " | ".join(parts)
        if "позиции по смете" in t:
            cols["key"] = col
        elif "по порядку" in t or t.startswith("№ п/п") or t.startswith("номер"):
            cols.setdefault("num", col)
        elif t.startswith("обоснование"):
            cols["code"] = col
        elif t.startswith("наименование"):
            cols["name"] = col
        elif t.startswith("единица"):
            cols["unit"] = col
        elif "всего с учетом коэффициентов" in t:
            cols["qty"] = col
        elif "базисном" in t:
            cols["price_base"] = col
        elif t.startswith("индекс") or "| индекс" in t:
            cols["index"] = col
        elif "на единицу измерения в текущем" in t:
            cols["price_cur"] = col
        elif "всего в текущем" in t:
            cols["total"] = col
        elif "коэффициент" in t:
            qty_cols.append(("coef", col))
        elif "на единицу" in t or t.startswith("количество"):
            qty_cols.append(("qty_base", col))
    # коэффициент к количеству — тот, что левее «всего с учетом коэффициентов»
    for kind, col in qty_cols:
        if kind == "coef" and col < cols.get("qty", 0):
            cols["coef"] = col
        if kind == "qty_base" and col < cols.get("qty", 0):
            cols.setdefault("qty_base", col)
    if "key" not in cols:          # в ЛСР «№ п/п» и есть номер позиции
        cols["key"] = cols.pop("num")
    need = {"key", "code", "name", "unit", "qty", "total"}
    miss = need - cols.keys()
    if miss:
        raise ValueError(f"Не распознаны колонки: {sorted(miss)}")
    return cols, num_row + 1


def _meta_ks2(ws) -> dict:
    meta = {}
    for row in ws.iter_rows(min_row=1, max_row=40):
        for c in row:
            t = _norm(txt(c.value))
            if t == "номер документа":
                meta["act_number"] = txt(ws.cell(c.row + 2, c.column).value)
            elif t == "дата составления":
                meta["act_date"] = txt(ws.cell(c.row + 2, c.column).value)
            elif t == "отчетный период":
                meta["period_from"] = txt(ws.cell(c.row + 2, c.column).value)
                meta["period_to"] = txt(ws.cell(c.row + 2, c.column + 1).value)
            elif t == "номер" and c.column > 1 and "договор" in _norm(txt(ws.cell(c.row, c.column - 1).value)):
                meta["contract_number"] = txt(ws.cell(c.row, c.column + 1).value)
            elif t == "дата" and "contract_number" in meta and "contract_date" not in meta:
                meta["contract_date"] = txt(ws.cell(c.row, c.column + 1).value)
            elif t in ("объект", "стройка") and t not in meta:
                vals = [txt(x.value) for x in row if x.column > c.column and txt(x.value)]
                if vals:
                    meta[t] = vals[0]
            elif t.startswith("заказчик"):
                vals = [txt(x.value) for x in row if x.column > c.column and txt(x.value)]
                meta.setdefault("customer", vals[0] if vals else "")
            elif t.startswith("подрядчик"):
                vals = [txt(x.value) for x in row if x.column > c.column and txt(x.value)]
                meta.setdefault("contractor", vals[0] if vals else "")
    return meta


def _meta_lsr(ws, hdr: int) -> tuple[dict, dict]:
    meta, summary = {}, {}
    for row in ws.iter_rows(min_row=1, max_row=hdr):
        for c in row:
            t = _norm(txt(c.value))
            if t.startswith("(наименование объекта"):
                prev = [txt(x.value) for x in ws[c.row - 1] if txt(x.value)]
                if prev:
                    meta["объект"] = prev[0]
            elif t.startswith("локальный сметный расчет"):
                meta["lsr_number"] = txt(c.value).split("№")[-1].strip()
            elif t.startswith("составлен(а) в текущем уровне цен"):
                vals = [txt(x.value) for x in row if x.column > c.column and txt(x.value)]
                meta["price_level"] = vals[0] if vals else ""
            elif t.startswith("сметная стоимость") or t in ("строительных работ", "монтажных работ",
                                                                 "оборудования", "прочих затрат"):
                vals = [num(x.value) for x in row if x.column > c.column and num(x.value) is not None]
                if vals:
                    summary[t.replace("сметная стоимость", "total").strip()] = vals[0]
    return meta, summary


def parse(path: str, kind: str, doc_id: Optional[str] = None) -> Document:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hdr = _find_header(ws)
    cols, start = _map_columns(ws, hdr)
    doc = Document(doc_id=doc_id or path.rsplit("/", 1)[-1], kind=kind, path=path, sheet=ws.title)
    if kind == "КС-2":
        doc.meta = _meta_ks2(ws)
    else:
        doc.meta, doc.header_summary = _meta_lsr(ws, hdr)

    g = lambda r, k: ws.cell(r, cols[k]).value if k in cols else None
    section = None
    open_pos: Optional[Position] = None
    in_final = False

    for r in range(start, ws.max_row + 1):
        key_raw, code, name = txt(g(r, "key")), txt(g(r, "code")), txt(g(r, "name"))
        first = key_raw or txt(ws.cell(r, 1).value)
        lname = _norm(name)

        # раздел
        if first.lower().startswith("раздел"):
            section = {"name": first, "total": None, "row": r, "keys": []}
            doc.sections.append(section)
            open_pos = None
            continue
        if lname.startswith("всего по разделу"):
            if section is not None:
                section["total"] = num(g(r, "total"))
            open_pos = None
            continue
        if lname.startswith("итоги по смете") or lname.startswith("всего прямые затраты (справочно)") and section is None:
            in_final = True

        # финальный блок: «Всего», коэффициент снижения, УСН/НДС, «ВСЕГО по …»
        if lname == "всего" or (in_final and lname == "всего"):
            doc.totals["vsego"] = (num(g(r, "total")), r)
            in_final = True
            continue
        if "коэффициента снижения" in lname or "коэффициент снижения" in lname:
            m = re.search(r"к\s*=\s*([\d.,]+)", name)
            doc.totals["reduction"] = (num(m.group(1)) if m else None, num(g(r, "total")), r, name)
            continue
        if "усн" in lname or "ндс" in lname:
            m = re.search(r"(\d+[.,]\d+|\d+)\s*$", name)
            doc.totals["tax_line"] = (num(m.group(1)) if m else None, num(g(r, "total")), r, name)
            continue
        if lname.startswith("всего по смете") or lname.startswith("всего по акту"):
            doc.totals["grand"] = (num(g(r, "total")), r)
            continue

        # позиция
        if KEY_RE.match(key_raw) and code and name and key_raw.count(".") <= 1:
            if key_raw in doc.positions or (open_pos and key_raw == open_pos.key):
                # повтор номера (5.1 → ресурс машины той же позиции)
                if open_pos:
                    open_pos.resources.append(Resource(r, code, name, txt(g(r, "unit")),
                                                       num(g(r, "qty")), num(g(r, "total"))))
                continue
            p = Position(doc_id=doc.doc_id, key=key_raw, row=r, code=code, name=name,
                         unit=txt(g(r, "unit")), qty_base=num(g(r, "qty_base")),
                         coef=num(g(r, "coef")), qty=num(g(r, "qty")),
                         section=section["name"] if section else "")
            if kind == "КС-2" and "num" in cols:
                p.act_num = txt(g(r, "num"))
            if open_pos is None:
                open_pos = p
            else:                                  # 28.1 внутри блока 28
                p.parent = open_pos.key
                p.unit_price = num(g(r, "price_cur"))
                p.total = num(g(r, "total"))
                open_pos.children.append(p.key)
            doc.positions[p.key] = p
            if section is not None:
                section["keys"].append(p.key)
            continue

        if open_pos is None:
            continue
        # строки внутри блока позиции
        if lname == "всего по позиции":
            open_pos.total = num(g(r, "total"))
            open_pos.unit_price = num(g(r, "price_cur"))
            open_pos.total_row = r
            open_pos = None
        elif lname == "итого прямые затраты":
            open_pos.direct = num(g(r, "total"))
        elif lname == "фот":
            open_pos.fot = num(g(r, "total"))
        elif lname.startswith("нр ") or lname.startswith("сп "):
            open_pos.overheads.append(Overhead(r, lname[:2].upper(), code, num(g(r, "qty")),
                                               num(g(r, "qty_base")), num(g(r, "coef")),
                                               num(g(r, "total"))))
        elif code:
            marker = first if first in MARKERS else (key_raw if key_raw in MARKERS else "")
            open_pos.resources.append(Resource(r, code, name, txt(g(r, "unit")),
                                               num(g(r, "qty")), num(g(r, "total")), marker))
    if not doc.positions:
        doc.parse_warnings.append("Не найдено ни одной позиции — проверьте формат выгрузки")
    return doc
