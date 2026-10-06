"""Step 8: клиентский Excel-отчёт + перевод пакета в SENT."""

import importlib
import os
import shutil
import sys
import tempfile
import unittest

import openpyxl


class DeliveryTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="delivery_test_")
        os.environ["MINIAPP_DATA_DIR"] = self.tmp_dir
        for mod in ("app.storage", "app.delivery_service", "app.client_report", "app.models"):
            sys.modules.pop(mod, None)
        self.storage = importlib.import_module("app.storage")
        self.models = importlib.import_module("app.models")
        self.svc = importlib.import_module("app.delivery_service")

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)
        os.environ.pop("MINIAPP_DATA_DIR", None)

    def _reviewed_package(self):
        pkg = self.models.Package.create(client_id="1", client_name='ООО "СЭМ"', object_title="Гидросталь")
        f_confirmed = self.models.Finding(id="f1", rule_id="S-02", status="WARNING",
                                           text="Объём превышен", decision="confirmed")
        f_edited = self.models.Finding(id="f2", rule_id="S-05", status="FAIL",
                                        text="Исходный текст", decision="edited",
                                        edited_text="Текст после правки оператора")
        f_rejected = self.models.Finding(id="f3", rule_id="H-03", status="WARNING",
                                          text="Это не должно попасть в отчёт", decision="rejected")
        pkg.set_findings([f_confirmed, f_edited, f_rejected])
        pkg.status = self.models.PackageStatus.REVIEWED.value
        self.storage.save_package(pkg)
        return pkg

    def test_prepare_client_report_includes_only_confirmed_and_edited(self):
        pkg = self._reviewed_package()
        updated = self.svc.prepare_client_report("1", pkg.id)

        self.assertIsNotNone(updated.client_report_path)
        path = self.storage.files_dir("1", pkg.id).parent / updated.client_report_path
        self.assertTrue(path.exists())

        wb = openpyxl.load_workbook(path)
        ws = wb["Замечания"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        data_rows = [r for r in rows if r[0] is not None and isinstance(r[0], int)]
        self.assertEqual(len(data_rows), 2)
        texts = [r[3] for r in data_rows]
        self.assertIn("Объём превышен", texts)
        self.assertIn("Текст после правки оператора", texts)
        self.assertNotIn("Это не должно попасть в отчёт", texts)
        self.assertNotIn("Исходный текст", texts)  # должна быть правка, не исходный текст

    def test_prepare_report_rejects_unfinished_package(self):
        pkg = self.models.Package.create(client_id="2")
        pkg.status = self.models.PackageStatus.READY_FOR_REVIEW.value
        self.storage.save_package(pkg)
        with self.assertRaises(self.svc.DeliveryError):
            self.svc.prepare_client_report("2", pkg.id)

    def test_mark_sent_requires_report_prepared_first(self):
        pkg = self._reviewed_package()
        with self.assertRaises(self.svc.DeliveryError):
            self.svc.mark_sent("1", pkg.id)

    def test_full_cycle_reviewed_to_sent(self):
        pkg = self._reviewed_package()
        self.svc.prepare_client_report("1", pkg.id)
        sent = self.svc.mark_sent("1", pkg.id)
        self.assertEqual(sent.status, self.models.PackageStatus.SENT.value)
        self.assertIsNotNone(sent.sent_at)

    def test_report_with_zero_confirmed_findings_still_builds_empty_sheet(self):
        pkg = self.models.Package.create(client_id="3")
        pkg.set_findings([self.models.Finding(id="f1", rule_id="S-01", status="OK",
                                                text="всё ок", decision="rejected")])
        pkg.status = self.models.PackageStatus.REVIEWED.value
        self.storage.save_package(pkg)

        updated = self.svc.prepare_client_report("3", pkg.id)
        path = self.storage.files_dir("3", pkg.id).parent / updated.client_report_path
        wb = openpyxl.load_workbook(path)
        ws = wb["Замечания"]
        data_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None and isinstance(r[0], int)]
        self.assertEqual(len(data_rows), 0)


if __name__ == "__main__":
    unittest.main()
