"""Step 8: клиентский Excel-отчёт — список подтверждённых оператором замечаний.

В отличие от внутреннего отчёта автопроверки (smeta_check.report, Step 5,
видит только оператор), сюда попадают ТОЛЬКО находки с решением
confirmed/edited — отклонённые (rejected) оператором намеренно не
показываются подрядчику: раз оператор решил, что замечание не по делу,
протаскивать его в отчёт клиента не нужно. Текст берётся из final_text()
(т.е. правка оператора, если она была, а не исходный текст автопроверки)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import FindingDecision, Package

FILL = {"FAIL": "F8CBAD", "WARNING": "FFE699", "NEEDS_CONTEXT": "BDD7EE", "INFO": "E7E6E6", "OK": "C6EFCE"}
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")

INCLUDED_DECISIONS = (FindingDecision.CONFIRMED.value, FindingDecision.EDITED.value)


def build_client_report(pkg: Package, out_path: Path) -> Path:
    included = [f for f in pkg.findings if f.decision in INCLUDED_DECISIONS]

    wb = Workbook()
    ws = wb.active
    ws.title = "Замечания"
    headers = ["№", "Правило", "Статус", "Замечание", "Документы"]
    widths = [5, 10, 15, 70, 34]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(1, i, h)
        c.font, c.alignment, c.border = FB, Alignment(wrap_text=True, vertical="center"), THIN
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(2, 1)

    for i, finding in enumerate(included, 1):
        row = i + 1
        docs = ", ".join(finding.source_docs) if finding.source_docs else ""
        vals = [i, finding.rule_id, finding.status, finding.final_text(), docs]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v)
            c.font, c.alignment, c.border = F, WRAP, THIN
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=FILL.get(finding.status, "FFFFFF"))

    if included:
        ws.auto_filter.ref = f"A1:E{len(included) + 1}"

    title_row = len(included) + 3
    title = ws.cell(title_row, 1,
                     f"Пакет: {pkg.object_title or pkg.client_name or pkg.client_id}. "
                     f"Подтверждённых замечаний: {len(included)}.")
    title.font = Font(name="Arial", size=9, italic=True, color="808080")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
